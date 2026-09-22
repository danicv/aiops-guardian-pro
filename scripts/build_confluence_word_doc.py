from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/deliverables/AIOps_Guardian_Confluence_Product_Guide.docx"
PRODUCT_DIAGRAM = ROOT / "docs/diagrams/end-to-end-product-architecture.png"
ENTERPRISE_DIAGRAM = ROOT / "docs/diagrams/end-to-end-enterprise-architecture.png"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None):
    result = doc.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(result.rows[0].cells[index], header, bold=True, color=(255, 255, 255))
        shade(result.rows[0].cells[index], "17365D")
    for row in rows:
        cells = result.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            if len(result.rows) % 2 == 0:
                shade(cells[index], "EAF2F8")
    if widths:
        for row in result.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph()
    return result


def bullet(doc, text, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.add_run(text)
    return paragraph


def numbered(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.add_run(text)
    return paragraph


def code_block(doc, text):
    paragraph = doc.add_paragraph(style="Code Block")
    paragraph.add_run(text)
    return paragraph


def add_mermaid(doc, title, source):
    doc.add_heading(title, level=3)
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("Mermaid source (paste into Confluence Mermaid macro)")
    run.italic = True
    code_block(doc, source)


def add_callout(doc, title, text, fill="D9EAF7"):
    result = doc.add_table(rows=1, cols=1)
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = result.cell(0, 0)
    shade(cell, fill)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(title + "\n")
    run.bold = True
    paragraph.add_run(text)
    doc.add_paragraph()


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    for name, size, color in [("Title", 30, "17365D"), ("Heading 1", 20, "17365D"), ("Heading 2", 15, "286090"), ("Heading 3", 11, "286090")]:
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    if "Code Block" not in [style.name for style in doc.styles]:
        style = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = doc.styles["Code Block"]
    style.font.name = "Aptos Mono"
    style.font.size = Pt(8)
    style.font.color.rgb = RGBColor(45, 45, 45)
    style.paragraph_format.left_indent = Inches(0.2)
    style.paragraph_format.space_before = Pt(4)
    style.paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("AIOps Guardian Pro | Confluence Product Guide | 22 September 2026").font.size = Pt(8)


def build():
    doc = Document()
    configure(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("AIOps Guardian Pro")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("SRE Incident Intelligence, Governed Remediation, and Operational Learning")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(40, 96, 144)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Confluence-ready product and end-user guide\nVersion 1.0 | 22 September 2026").italic = True
    doc.add_picture(str(PRODUCT_DIAGRAM), width=Inches(6.65))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Figure 1. End-to-end product architecture.").alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_callout(doc, "One-line value proposition", "AIOps Guardian turns operational signals into explainable, evidence-backed, and governed SRE actions.")

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph("AIOps Guardian Pro is an SRE control plane for detecting, investigating, explaining, and safely remediating production issues. It combines a React operations console, FastAPI services, a LangGraph orchestrator, specialized diagnostic agents, live observability integrations, deterministic guardrails, human approvals, MCP tool execution, verification, and an operational audit trail.")
    doc.add_paragraph("An end user does not need to understand Kubernetes, Prometheus, tracing, MCP, or agent orchestration. They can ask a question such as:")
    add_callout(doc, "Example user question", "Why is checkout-api failing after today's production deployment?")
    doc.add_paragraph("Guardian collects evidence, correlates signals, explains the likely cause, proposes a practical response, and pauses risky production actions until an authorized approver decides.")
    table(doc, ["Capability", "What the user gets"], [
        ("Investigate", "Plain-language incident analysis with evidence and confidence."),
        ("Observe", "Metrics, logs, Kubernetes health, deployment context, and telemetry quality."),
        ("Decide", "Recommended containment and remediation with policy reasoning."),
        ("Control", "Human approval for risky actions and single-use approval decisions."),
        ("Learn", "Persisted incidents, audit events, reports, and historical comparisons."),
    ], [1.3, 5.3])

    doc.add_heading("2. Product Scope", level=1)
    table(doc, ["Area", "Included behavior"], [
        ("Incident intelligence", "Correlates metrics, logs, Kubernetes events, pipeline changes, runbooks, and incident history."),
        ("Release intelligence", "Identifies risky releases, missing evidence, configuration regressions, and promotion concerns."),
        ("SRE guidance", "Produces advanced metrics, PromQL examples, containment steps, remediation steps, and verification criteria."),
        ("Governance", "Separates read-only analysis from controlled actions and records the decision trail."),
        ("Integrations", "Supports live Prometheus and Loki connections with validated URLs and environment-backed credentials."),
        ("Deployment", "Runs locally on kind Kubernetes and is structured for GKE/GCP deployment."),
    ], [1.6, 5.0])

    doc.add_heading("3. End-User Experience", level=1)
    doc.add_heading("Primary workflow", level=2)
    for item in [
        "Open the Dashboard to see current incident, approval, release-risk, and platform-health signals.",
        "Open Ask Guardian and describe the problem in natural language.",
        "Review the root cause, confidence, evidence sources, agent trace, metrics, and recommended action.",
        "Follow the SRE action plan: contain, remediate, and verify.",
        "If the action is high risk, review it in Approvals and approve or reject it.",
        "Use Investigations and Reports to review history, telemetry quality, and recovery evidence.",
    ]:
        numbered(doc, item)
    table(doc, ["Screen", "Purpose", "Typical user action"], [
        ("Dashboard", "Operational overview", "Identify what needs attention first."),
        ("Ask Guardian", "Investigation and release-risk analysis", "Ask a question and review the evidence-backed response."),
        ("Live Reports", "Telemetry trends and data quality", "Inspect request rate, errors, latency, burn rate, restarts, and logs."),
        ("Simulator", "Training and demo data", "Generate deterministic incidents and exercise approvals."),
        ("Investigations", "Incident history", "Review prior investigations and root causes."),
        ("Pipelines", "Release intelligence", "Compare pipeline status, version, and risk."),
        ("Approvals", "Human-in-the-loop control", "Approve or reject pending production actions."),
        ("Integrations", "Data-source management", "Test, save, and activate Prometheus/Loki sources."),
        ("Architecture", "Platform explanation", "Explain how agents, MCP, guardrails, and observability fit together."),
    ], [1.3, 2.2, 3.1])

    doc.add_heading("4. Incident Investigation Example", level=1)
    doc.add_paragraph("The reference scenario is an unsafe checkout-api release that changes the container memory limit from 1Gi to 128Mi.")
    table(doc, ["Signal", "Before", "After", "Interpretation"], [
        ("HTTP 5xx rate", "0.2%", "6.8%", "Customer-visible failures."),
        ("p95 latency", "180 ms", "690 ms", "Severe performance degradation."),
        ("Pod restarts", "0", "9", "Runtime instability."),
        ("Error-budget burn", "1x", "34x", "SLO risk is accelerating."),
        ("Termination reason", "Healthy", "OOMKilled", "Memory limit is below demand."),
    ], [1.6, 1.0, 1.0, 3.0])
    add_callout(doc, "Example conclusion", "The latest deployment reduced the memory limit below application demand, causing OOMKilled terminations, pod restarts, HTTP 5xx errors, and latency degradation.", "FCE4D6")
    doc.add_heading("Recommended operating response", level=2)
    for item in [
        "Contain: pause promotion and route traffic to the last known-good version.",
        "Remediate: restore the validated memory limit or roll back to v1.9.0.",
        "Verify: require stable 5xx, p95 latency, restart, and error-budget metrics for 15 minutes.",
    ]:
        numbered(doc, item)

    doc.add_heading("5. Architecture and Control Boundaries", level=1)
    doc.add_picture(str(ENTERPRISE_DIAGRAM), width=Inches(6.65))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Figure 2. Enterprise architecture and integration boundaries.").alignment = WD_ALIGN_PARAGRAPH.CENTER
    table(doc, ["Component", "Responsibility"], [
        ("React/Vite UI", "User workflow, investigation results, reports, simulator, and approvals."),
        ("FastAPI", "API boundary, authentication, persistence orchestration, and route-level authorization."),
        ("LangGraph orchestrator", "Intent, plan, shared state, agent routing, validation, guardrails, and completion."),
        ("Specialized agents", "Pipeline, Kubernetes, monitoring, change, RAG, history, RCA, release risk, and verification."),
        ("MCP tool plane", "Governed access to Kubernetes, CI/CD, metrics, logs, traces, notifications, and actions."),
        ("Guardrails", "Deterministically allow, block, or require approval for operational actions."),
        ("PostgreSQL and FAISS", "Structured audit/incident memory and operational knowledge retrieval."),
        ("OpenTelemetry", "Trace agent execution, tool calls, approvals, remediation, and verification."),
    ], [1.7, 4.9])
    add_mermaid(doc, "Editable architecture diagram", """flowchart TD\n  U[User / Alert / Pipeline Event] --> UI[React Operations UI]\n  UI --> API[FastAPI API]\n  API --> O[LangGraph Orchestrator]\n  O --> A[Specialized Agents]\n  A --> E[Evidence Correlation and RCA]\n  E --> V[Validation and Adaptive Routing]\n  V --> G[Guardrail Policy]\n  G -->|Allow| M[Action MCP]\n  G -->|Require approval| H[Human Approval]\n  H --> M\n  M --> X[Verification]\n  X --> DB[(PostgreSQL Audit and Incident Memory)]\n  A --> OBS[Prometheus / Loki / Tempo / OTel]\n  A --> RAG[FAISS and Runbooks]""")

    doc.add_heading("6. Safety, Identity, and Authorization", level=1)
    doc.add_paragraph("Guardian separates evidence collection from action execution. Read-only investigation can be performed by a viewer. Production actions require an approver role, and approval decisions are single-use to prevent replay.")
    table(doc, ["Role", "Allowed behavior"], [
        ("Viewer", "Read dashboards, reports, investigations, pipeline data, integrations, and run conversations."),
        ("Approver", "Everything a viewer can do, plus approve/reject actions, activate integrations, add integrations, and send operational notifications."),
        ("Admin", "Everything above, subject to organization policy and deployment controls."),
    ], [1.4, 5.2])
    bullet(doc, "Production Kubernetes configuration enables authentication by default.")
    bullet(doc, "Local demo mode explicitly disables authentication for disposable development use.")
    bullet(doc, "Integration URLs reject metadata, private, loopback, reserved, multicast, and link-local addresses.")
    bullet(doc, "Credentials are referenced through environment-backed secrets rather than saved as raw integration values.")
    bullet(doc, "PostgreSQL deployments use Alembic migrations; SQLite table creation remains limited to local demo/test use.")

    doc.add_heading("7. Observability and SRE Metrics", level=1)
    table(doc, ["Metric family", "Purpose", "Example"], [
        ("RED metrics", "Request health", "Rate, HTTP 5xx ratio, duration histogram."),
        ("Saturation", "Resource pressure", "Memory working set, memory limit, CPU throttling."),
        ("Reliability", "SLO health", "Error-budget burn rate and scrape freshness."),
        ("Kubernetes", "Runtime state", "Restarts, OOMKilled, desired versus healthy replicas."),
        ("Trace correlation", "Causality", "W3C trace context and deployment.version."),
    ], [1.5, 2.4, 2.7])
    doc.add_paragraph("Guardian marks unavailable data as unknown or insufficient rather than treating missing telemetry as healthy. Each report exposes scope, timestamp, source, query, status, and data gaps.")

    doc.add_heading("8. Local Deployment", level=1)
    code_block(doc, "make local-build\nkind load docker-image guardian-backend:reports-v1 guardian-mcp:local guardian-frontend:reports-v1 --name guardian-local\nmake local-deploy")
    table(doc, ["Service", "Local access"], [
        ("Frontend", "http://localhost:30081"),
        ("Backend Swagger", "http://localhost:30800/docs"),
        ("Prometheus", "ClusterIP; port-forward when needed"),
        ("Loki", "ClusterIP; used by Live Reports"),
        ("PostgreSQL", "Internal ClusterIP service with persistent storage"),
    ], [2.0, 4.6])
    doc.add_paragraph("The local kind deployment includes the frontend, backend, MCP server, checkout-api telemetry target, Prometheus, Loki, PostgreSQL, OpenTelemetry sidecars, and a persistent Prometheus volume.")

    doc.add_heading("9. GCP / GKE Deployment Path", level=1)
    for item in [
        "Bootstrap Terraform state, Workload Identity Federation, and the deployment service account.",
        "Create Secret Manager versions for database, OpenAI, SMTP, and authentication secrets.",
        "Build commit-addressed images and push them to Artifact Registry.",
        "Apply Terraform to provision GKE, networking, IAM, secrets, and the Helm release.",
        "Expose the frontend and /api through a managed ingress with TLS and Cloud Armor.",
        "Place the application behind Google Cloud IAP or an organization-managed OIDC gateway.",
        "Use Cloud SQL for production PostgreSQL and persistent storage for telemetry systems.",
    ]:
        numbered(doc, item)
    add_callout(doc, "GCP production gate", "Do not expose the API publicly until identity-aware access, Secret Manager, managed TLS, network policy, database migrations, and action authorization are configured.", "FFF2CC")

    doc.add_heading("10. Confluence Publishing Guide", level=1)
    doc.add_paragraph("Recommended Confluence parent page: AIOps Guardian Pro")
    table(doc, ["Child page", "Content"], [
        ("Product Overview", "Executive summary, value proposition, user journey, and scope."),
        ("End User Guide", "Dashboard, Ask Guardian, Reports, Simulator, Investigations, Pipelines, Approvals, and Integrations."),
        ("SRE Runbook", "Incident workflow, metrics, containment, remediation, verification, and escalation."),
        ("Architecture", "Architecture diagram, responsibility boundaries, MCP tool plane, and data flow."),
        ("Security Model", "Identity, roles, guardrails, secrets, SSRF controls, and audit requirements."),
        ("Local Kubernetes", "kind setup, image loading, port forwarding, simulator, and telemetry."),
        ("GCP / GKE", "Terraform bootstrap, Secret Manager, Artifact Registry, Workload Identity, and ingress."),
        ("API Reference", "Swagger URL, route groups, authentication, and example payloads."),
    ], [2.0, 4.6])
    doc.add_paragraph("For editable diagrams, copy the Mermaid source blocks from the appendix into a Confluence Mermaid macro. Keep the PNG architecture diagrams as the executive-friendly visual fallback.")

    doc.add_heading("11. Editable Diagram Sources", level=1)
    add_mermaid(doc, "Incident response sequence", """sequenceDiagram\n  participant User\n  participant UI as Guardian UI\n  participant API as FastAPI\n  participant O as Orchestrator\n  participant Obs as Metrics/Logs/K8s\n  participant G as Guardrail\n  participant A as Approver\n  participant MCP as Action MCP\n  participant V as Verification\n  User->>UI: Ask incident question\n  UI->>API: Investigation request\n  API->>O: Create investigation\n  O->>Obs: Collect scoped evidence\n  Obs-->>O: Metrics, logs, events, changes\n  O->>O: Correlate and validate RCA\n  O->>G: Evaluate recommended action\n  G-->>O: Require approval\n  O-->>UI: RCA, metrics, action plan, approval\n  A->>API: Approve or reject\n  API->>MCP: Execute approved action\n  MCP-->>API: Action result\n  API->>V: Verify recovery\n  V-->>UI: Recovery status and audit trail""")
    add_mermaid(doc, "Deployment topology", """flowchart LR\n  subgraph GKE[GKE / Local kind]\n    FE[Frontend] --> BE[FastAPI Backend]\n    BE --> DB[(PostgreSQL)]\n    BE --> MCP[MCP Server]\n    BE --> P[Prometheus]\n    BE --> L[Loki]\n    BE --> O[OTel Sidecar]\n    P --> PVC[(Persistent Prometheus Volume)]\n  end\n  O --> OTG[OTel Gateway]\n  BE --> SM[Secret Manager / Kubernetes Secrets]\n  User --> IAP[OIDC or Google IAP]\n  IAP --> FE""")

    doc.add_heading("12. Source References", level=1)
    for reference in [
        "README.md - product scope and quick start",
        "docs/architecture-one-page.md - responsibility boundaries and closed-loop flow",
        "docs/security.md - identity, authorization, secrets, network, AI controls, and audit",
        "docs/live-integrations.md - live telemetry and controlled traffic scenario",
        "docs/deployment-gcp.md - Terraform, GKE, Secret Manager, and deployment workflow",
        "docs/observability.md - OTel, Prometheus, Loki, Tempo, and recommended spans",
        "demo/README.md - unsafe memory release incident scenario",
    ]:
        bullet(doc, reference)

    doc.add_heading("Document Maintenance", level=1)
    doc.add_paragraph("Owner: Platform Engineering / SRE\nReview cadence: Each release or architecture change\nLast validated: 22 September 2026\nRepository: AIOps Guardian Pro")
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
